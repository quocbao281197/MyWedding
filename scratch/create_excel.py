import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import re

txt_path = r"D:\Working_Bao\Đám cưới\Hình Cưới\Final filter\Ghi_chu_sua_anh.txt"
excel_path = r"D:\Working_Bao\Đám cưới\Hình Cưới\Final filter\Danh_sach_hinh_can_sua.xlsx"

with open(txt_path, "r", encoding="utf-8") as f:
    lines = f.readlines()

photos_data = []

current_code = None
current_content = []
current_notes = []

for line in lines:
    line_str = line.strip()
    
    # Detect photo header like [01] TLAT1479.JPG or [35] VHH_8234.JPG
    match = re.search(r"\[\d+\]\s+([\w\d\._]+)", line_str)
    if match:
        # Save previous photo if exists
        if current_code:
            photos_data.append((
                current_code, 
                "\n".join(current_content).strip(), 
                "\n".join(current_notes).strip()
            ))
        current_code = match.group(1).strip()
        current_content = []
        current_notes = []
        continue
    
    if not current_code:
        continue
        
    # Ignore delimiter lines
    if line_str.startswith("----------------") or line_str.startswith("================"):
        continue
        
    cleaned_line = line_str.replace("**", "")
    
    # Classify into Note or Content
    if "Xoay ảnh:" in cleaned_line or "Xoay & Bố cục:" in cleaned_line:
        note_text = cleaned_line.replace("• Xoay ảnh:", "").replace("• Xoay & Bố cục:", "").replace("•", "").strip()
        current_notes.append(note_text)
    else:
        if cleaned_line:
            current_content.append(cleaned_line)

# Save the last photo
if current_code:
    photos_data.append((
        current_code, 
        "\n".join(current_content).strip(), 
        "\n".join(current_notes).strip()
    ))

print(f"Total photos parsed: {len(photos_data)}")
# Build Excel File

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Ghi Chú Retouch"
ws.views.sheetView[0].showGridLines = True

# Title Row
ws.merge_cells("A1:D1")
title_cell = ws["A1"]
title_cell.value = "DANH SÁCH CHI TIẾT CÁC ĐIỂM CẦN CHỈNH SỬA HÌNH CƯỚI"
title_cell.font = Font(name="Segoe UI", size=14, bold=True, color="1F4E79")
title_cell.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 36

# Headers
headers = ["STT", "Mã hình chỉnh sửa", "Nội dung chỉnh sửa", "Note"]
ws.row_dimensions[2].height = 26

header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

for col_idx, h in enumerate(headers, 1):
    cell = ws.cell(row=2, column=col_idx, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border

# Fill Rows
zebra_fill = PatternFill(start_color="F9FAFC", end_color="F9FAFC", fill_type="solid")

for idx, (code, content, note) in enumerate(photos_data, start=1):
    r = idx + 2
    ws.cell(row=r, column=1, value=idx)
    ws.cell(row=r, column=2, value=code)
    ws.cell(row=r, column=3, value=content)
    ws.cell(row=r, column=4, value=note if note else "Khung đứng chuẩn")
    
    c1 = ws.cell(row=r, column=1)
    c2 = ws.cell(row=r, column=2)
    c3 = ws.cell(row=r, column=3)
    c4 = ws.cell(row=r, column=4)
    
    c1.alignment = Alignment(horizontal="center", vertical="top")
    c2.alignment = Alignment(horizontal="center", vertical="top")
    c3.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    c4.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    
    font_body = Font(name="Segoe UI", size=10)
    c1.font = font_body
    c2.font = Font(name="Segoe UI", size=10, bold=True, color="002060")
    c3.font = font_body
    c4.font = Font(name="Segoe UI", size=10, italic=True, color="C00000" if note else "555555")
    
    for c in [c1, c2, c3, c4]:
        c.border = thin_border
        if idx % 2 == 0:
            c.fill = zebra_fill

ws.column_dimensions['A'].width = 8
ws.column_dimensions['B'].width = 22
ws.column_dimensions['C'].width = 85
ws.column_dimensions['D'].width = 35

wb.save(excel_path)
print("Excel rebuilt successfully!")
