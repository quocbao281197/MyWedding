import openpyxl

file_path = r'D:\Working_Bao\Đám cưới\Hình Cưới\THÔNG TIN CHỈNH SỬA HÌNH ẢNH.xlsx'
wb = openpyxl.load_workbook(file_path)
ws = wb.active

with open(r'd:\Working_Bao\MyRepositories\Wedding\MyWedding\scratch\user_excel_dump.txt', 'w', encoding='utf-8') as f:
    f.write(f"Sheet: {ws.title}, max_row: {ws.max_row}, max_col: {ws.max_column}\n")
    for i in range(1, ws.max_row + 1):
        row_vals = [str(ws.cell(row=i, column=j).value) if ws.cell(row=i, column=j).value is not None else "" for j in range(1, ws.max_column + 1)]
        f.write(f"Row {i}: {' | '.join(row_vals)}\n")

print("Dumped user excel successfully!")
